"""Markdown → DOCX 转换（支持标题/表格/列表/粗体/内联代码/引用/代码块）。

用法：python scripts/md2docx.py <输入.md> <输出.docx>
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt


def add_runs(paragraph: object, text: str) -> None:
    """处理 **粗体** 与 `内联代码`。"""
    for part in re.split(r"(\*\*.*?\*\*|`[^`]*`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def is_sep(row: str) -> bool:
    return re.match(r"^\|[\s:|-]+\|$", row) is not None


def parse_row(row: str) -> list[str]:
    return [c.strip() for c in row.strip().strip("|").split("|")]


def convert(md_path: str, docx_path: str) -> None:
    doc = Document()
    lines = Path(md_path).read_text(encoding="utf-8").split("\n")
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                r = lines[i].strip()
                if not is_sep(r):
                    rows.append(parse_row(r))
                i += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Light Grid Accent 1"
                for ri, row in enumerate(rows):
                    for ci in range(len(rows[0])):
                        t.rows[ri].cells[ci].text = row[ci] if ci < len(row) else ""
            i -= 1
        elif s.startswith("- "):
            doc.add_paragraph(s[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s", s):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", s), style="List Number")
        elif s.startswith(">"):
            doc.add_paragraph(s.lstrip("> ").strip())
        elif s in ("", "---"):
            pass
        else:
            p = doc.add_paragraph()
            add_runs(p, s)
        i += 1
    doc.save(docx_path)
    print(f"已生成 {docx_path}")


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
